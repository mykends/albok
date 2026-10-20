import shutil
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from albok_api.documents import convert_to_pdf, section_pdf
from albok_api.importers import BookImportError, read_pdf, suggest_sections
from albok_api.models import BookFile
from albok_api.pdfgen import image_to_pdf, text_to_pdf


def make_pdf(pages: list[str]) -> bytes:
    """Build a minimal text PDF (one line per page) without extra dependencies."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", ""]
    kids = []
    for text in pages:
        page_id, content_id = len(objects) + 1, len(objects) + 2
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"

    out = BytesIO(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n{body}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def ranges(sections):
    return [(s.title, s.page_start, s.page_end) for s in sections]


def test_outline_becomes_page_ranges_with_front_matter():
    pages = ["Naslovnica", "Sadržaj", "1 Uvod u", "više", "2 Dalje"]
    sections = suggest_sections(pages, [("Prvo", 2), ("Drugo", 4)], fallback_title="x")
    assert ranges(sections) == [("Uvod", 1, 2), ("Prvo", 3, 4), ("Drugo", 5, 5)]
    assert "Naslovnica" in sections[0].text


def test_text_headings_share_a_page_when_not_at_the_top():
    pages = [
        "Poglavlje 1 Stanice\nStanice su male.",
        "Još o stanicama.\nNešto.\nI još.\nPoglavlje 2 Tkiva\nTkiva su skupine stanica.",
        "Poglavlje 3 Organi\nOrgani.",
    ]
    assert ranges(suggest_sections(pages, [], fallback_title="Bio")) == [
        ("Poglavlje 1 Stanice", 1, 2),
        ("Poglavlje 2 Tkiva", 2, 2),
        ("Poglavlje 3 Organi", 3, 3),
    ]


def test_numbered_headings_split_on_the_shallowest_repeated_level():
    pages = ["2.1 Linearne jednadžbe\ntekst", "2.1.1 Primjer\ntekst", "2.2 Nejednadžbe\ntekst"]
    assert [s.title for s in suggest_sections(pages, [], fallback_title="x")] == ["2.1 Linearne jednadžbe", "2.2 Nejednadžbe"]


def test_numbered_list_items_are_not_headings():
    pages = ["Koraci:\n1. Pomiješaj otopinu\n2. Zagrij je\n3. Zapiši rezultat"]
    assert ranges(suggest_sections(pages, [], fallback_title="Vježba")) == [("Vježba", 1, 1)]


def test_known_word_headings_are_located_in_pages():
    pages = ["Egipat\nNil.", "Rim\nGrad-država.", "Još Rima."]
    sections = suggest_sections(pages, [], fallback_title="x", known_headings=["Egipat", "Rim"])
    assert ranges(sections) == [("Egipat", 1, 1), ("Rim", 2, 3)]


def test_slides_are_one_section_and_long_files_are_chunked():
    assert ranges(suggest_sections(["a", "b", "c"], [("x", 0), ("y", 1)], fallback_title="Prezentacija", slides=True)) == [
        ("Prezentacija", 1, 3)
    ]
    chunks = suggest_sections([""] * 25, [], fallback_title="Sken")
    assert ranges(chunks) == [("Stranice 1–10", 1, 10), ("Stranice 11–20", 11, 20), ("Stranice 21–25", 21, 25)]


def test_scanned_pdf_without_text_still_works(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(595, 842)
    writer.write(str(path))
    pages, outline, _ = read_pdf(path)
    assert len(pages) == 3 and outline == []
    assert ranges(suggest_sections(pages, outline, fallback_title="Sken")) == [("Sken", 1, 3)]


@pytest.mark.parametrize(("data", "message"), [(b"not a pdf", "ne može pročitati")])
def test_broken_pdf_raises_readable_error(tmp_path, data, message):
    path = tmp_path / "x.pdf"
    path.write_bytes(data)
    with pytest.raises(BookImportError, match=message):
        read_pdf(path)


def test_markdown_renders_with_bookmarks_and_croatian_letters(tmp_path):
    out = tmp_path / "md.pdf"
    text_to_pdf("# Knjiga\n\n## Čvorovi\n\nŽuta šuma, đak i ćup.\n\n## Šume\n\n```\nkôd č\n```\n", "Knjiga", out)
    pages, outline, title = read_pdf(out)
    assert title == "Knjiga"
    assert [name for name, _ in outline] == ["Čvorovi", "Šume"]
    assert "Žuta šuma, đak i ćup." in pages[0]


def test_image_becomes_a_one_page_pdf(tmp_path):
    from PIL import Image

    image = tmp_path / "ploca.webp"
    Image.new("RGB", (800, 600), "white").save(image)
    out = tmp_path / "ploca.pdf"
    image_to_pdf(image, out)
    page = PdfReader(str(out)).pages[0]
    assert abs(float(page.mediabox.width) / float(page.mediabox.height) - 800 / 600) < 0.01


def test_section_pdf_contains_only_its_pages(tmp_path, data_dir):
    pdf = data_dir / "files" / "books" / "1" / "1" / "document.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(make_pdf(["jedan", "dva", "tri", "cetiri"]))
    book_file = BookFile(id=1, book_id=1, original_filename="x.pdf", original_path="", pdf_path="books/1/1/document.pdf")
    reader = PdfReader(str(section_pdf(book_file, 2, 3)))
    assert [page.extract_text() for page in reader.pages] == ["dva", "tri"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice is not installed")
def test_word_document_converts_with_heading_bookmarks(tmp_path, data_dir):
    document = Document()
    document.add_heading("Svjetska povijest", level=0)
    document.add_heading("Stari Egipat", level=1)
    document.add_paragraph("Nil je oblikovao život u Egiptu.")
    document.add_page_break()
    document.add_heading("Stari Rim", level=1)
    document.add_paragraph("Rim je nastao kao grad-država.")
    original = tmp_path / "povijest.docx"
    document.save(original)

    out = tmp_path / "povijest.pdf"
    convert_to_pdf(original, "povijest.docx", out)
    pages, outline, _ = read_pdf(out)
    sections = suggest_sections(pages, outline, fallback_title="povijest")
    assert [s.title for s in sections][-2:] == ["Stari Egipat", "Stari Rim"]
    assert sections[-1].page_start == 2
